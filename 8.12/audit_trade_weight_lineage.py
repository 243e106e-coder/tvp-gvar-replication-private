#!/usr/bin/env python3
"""Read-only inventory of trade weight inputs."""
import hashlib,json,pathlib
from openpyxl import load_workbook
files=['.github/workflows/Trade_Weights_14_Economies_Stage1_2000_2014.xlsx','8.12/Trade_Weights_14_Economies_2000_2012(2).xlsx','8.12/Trade_Weights_14_Economies_2000_2014.csv','.github/workflows/Trade_Weights_14_Economies_2000_2014.csv']
out=pathlib.Path('results/trade_weight_lineage');out.mkdir(parents=True,exist_ok=True)
report=[]
for name in files:
 p=pathlib.Path(name)
 if not p.is_file():raise FileNotFoundError(name)
 item={'file':name,'bytes':p.stat().st_size,'sha256':hashlib.sha256(p.read_bytes()).hexdigest()}
 if p.suffix=='.xlsx':
  book=load_workbook(p,read_only=True,data_only=False)
  item['sheets']=[{'name':s.title,'rows':s.max_row,'cols':s.max_column} for s in book]
  book.close()
 report.append(item)
(out/'inventory.json').write_text(json.dumps(report,indent=2,ensure_ascii=False)+'\n',encoding='utf-8')
(out/'LIMITATIONS.txt').write_text('This is only an inventory. The documented 2013-2014 mirror completion remains unverified at pair-year level. Do not substitute zeros, mirrors or interpolations for missing data.\n')
print('Inventory completed',len(report))
