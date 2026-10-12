#!/usr/bin/env python3
"""Read-only workbook cell and formula provenance audit. No data imputation."""
import csv,json,pathlib,re
from openpyxl import load_workbook
P=[
".github/workflows/Trade_Weights_14_Economies_Stage1_2000_2014.xlsx",
"8.12/Trade_Weights_14_Economies_2000_2012(2).xlsx"]
out=pathlib.Path("results/trade_workbook_cells");out.mkdir(parents=True,exist_ok=True)
metadata=[];cells=[];signals=[]
terms=re.compile(r"2013|2014|mirror|imput|fill|source|comtrade|dots|exports?|imports?|euro|aggregate|missing|缺失|补|来源",re.I)
for filename in P:
 wb=load_workbook(filename,read_only=False,data_only=False)
 for ws in wb.worksheets:
  metadata.append(dict(file=filename,sheet=ws.title,rows=ws.max_row,columns=ws.max_column))
  for row in ws.iter_rows():
   for c in row:
    if c.value is None:continue
    val=str(c.value)
    cells.append(dict(file=filename,sheet=ws.title,cell=c.coordinate,value=val[:2500],is_formula=val.startswith("=")))
    if terms.search(val) or (c.comment and terms.search(c.comment.text)):
     signals.append(dict(file=filename,sheet=ws.title,cell=c.coordinate,text=val[:2500],comment=(c.comment.text[:1000] if c.comment else "")))
 wb.close()
with (out/"workbook_cells.csv").open("w",newline="",encoding="utf-8") as f:
 w=csv.DictWriter(f,fieldnames=("file","sheet","cell","value","is_formula"));w.writeheader();w.writerows(cells)
with (out/"provenance_signals.csv").open("w",newline="",encoding="utf-8") as f:
 w=csv.DictWriter(f,fieldnames=("file","sheet","cell","text","comment"));w.writeheader();w.writerows(signals)
(out/"sheet_dimensions.json").write_text(json.dumps(metadata,indent=2)+"\n",encoding="utf-8")
(out/"LIMITATIONS.md").write_text("# Provenance limitations\n\nThis exports all nonempty workbook cells and textual provenance signals without changing any workbook. A reported 2013/2014 label or a spreadsheet formula does NOT prove underlying trade observation or official origin. Pair-year mirror/NA lineage is UNVERIFIED unless explicitly recorded in source materials. Do not substitute mirrored, zero, or interpolated data for missing records. To certify the final 14x14 W, obtain original reporter-partner-year trade values and documented aggregation including EA.\n",encoding="utf-8")
print("Audited",len(metadata),"worksheets;",len(cells),"cells;",len(signals),"keyword matches")
